"""Xuất bộ test case (JSON theo references/schema-json.md) ra Excel theo mẫu Trí Nam.

Cách dùng:
    python xuat_excel.py <dau_vao.json> <dau_ra.xlsx>

Script kiểm tra dữ liệu trước khi xuất. Có lỗi → in danh sách lỗi, thoát mã 1, không ghi file.
"""
import json
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LOAI_TC = {"Chức năng", "Giao diện", "Phân quyền", "Biên", "Ngoại lệ", "API", "Hiệu năng", "ATTT"}
KY_THUAT = {"PVTĐ", "GTB", "BQĐ", "CTT", "PQ", "LUC", "ĐL"}
UU_TIEN = {"Cao", "Trung bình", "Thấp"}
LOAI_DX = {"Mơ hồ", "Thiếu", "Mâu thuẫn", "Không kiểm chứng được"}
NHOM_YC = {"Use case", "Quy tắc nghiệp vụ", "Phi chức năng", "Khác"}
NHOM_TC = {"Happy case", "Chức năng", "UI/UX", "Hiệu năng"}  # nhóm kịch bản theo quy trình Trí Nam
CHE_DO = {"Đầy đủ", "Gọn"}
TT_DX_CHO, TT_DX_DA_TRA_LOI, TT_DX_DA_CAP_NHAT = "Chờ trả lời", "Đã trả lời", "Đã cập nhật TC"
TRANG_THAI_DX = {TT_DX_CHO, TT_DX_DA_TRA_LOI, TT_DX_DA_CAP_NHAT}

TRUONG_TC_BAT_BUOC = ["ma_tc", "ma_yeu_cau", "chuc_nang", "tieu_de", "loai", "ky_thuat", "uu_tien",
                      "dieu_kien_truoc", "buoc_thuc_hien", "ket_qua_mong_doi"]

# Tên cột dùng chung với doc_tra_loi.py – đổi ở đây thì hai script cùng đổi.
COT_DX_MA, COT_DX_TRA_LOI, COT_DX_TRANG_THAI = "Mã", "Trả lời của BA/KH", "Trạng thái DX"
COT_TC_MA = "Mã TC"
COT_TC_THUC_THI = {  # cột tester điền khi chạy test → khoá trong JSON
    "Kết quả thực tế": "ket_qua_thuc_te",
    "Trạng thái": "trang_thai",
    "Mã lỗi": "ma_loi",
    "Người thực hiện / Ngày": "nguoi_thuc_hien",
}

MAU_TIEU_DE = PatternFill("solid", fgColor="1F4E78")
MAU_CANH_BAO = PatternFill("solid", fgColor="FCE4D6")
MAU_THIEU = PatternFill("solid", fgColor="F8CBAD")
MAU_DA_XONG = PatternFill("solid", fgColor="E2EFDA")
VIEN = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def kiem_tra(du_lieu):
    loi = []
    tt = du_lieu.get("thong_tin", {})
    yeu_cau = du_lieu.get("yeu_cau", [])
    dx = du_lieu.get("diem_can_xac_nhan", [])
    tcs = du_lieu.get("test_cases", [])

    che_do = tt.get("che_do", "Đầy đủ")
    if che_do not in CHE_DO:
        loi.append(f"thong_tin.che_do không hợp lệ '{che_do}'")

    ma_yc = [y.get("ma") for y in yeu_cau]
    for ma, n in Counter(ma_yc).items():
        if n > 1:
            loi.append(f"Mã yêu cầu trùng: {ma}")
    for y in yeu_cau:
        if y.get("nhom") not in NHOM_YC:
            loi.append(f"{y.get('ma')}: nhóm yêu cầu không hợp lệ '{y.get('nhom')}'")

    ma_dx = [d.get("ma") for d in dx]
    for ma, n in Counter(ma_dx).items():
        if n > 1:
            loi.append(f"Mã DX trùng: {ma}")
    dx_da_chot = set()
    for d in dx:
        if d.get("loai") not in LOAI_DX:
            loi.append(f"{d.get('ma')}: loại không hợp lệ '{d.get('loai')}'")
        if d.get("muc_anh_huong") not in UU_TIEN:
            loi.append(f"{d.get('ma')}: mức ảnh hưởng không hợp lệ '{d.get('muc_anh_huong')}'")
        for truong in ("vi_tri", "noi_dung", "cau_hoi"):
            if not d.get(truong):
                loi.append(f"{d.get('ma')}: thiếu '{truong}'")
        trang_thai = d.get("trang_thai", TT_DX_CHO)
        if trang_thai not in TRANG_THAI_DX:
            loi.append(f"{d.get('ma')}: trạng thái DX không hợp lệ '{trang_thai}'")
        if trang_thai != TT_DX_CHO:
            dx_da_chot.add(d.get("ma"))
            if not d.get("tra_loi"):
                loi.append(f"{d.get('ma')}: trạng thái '{trang_thai}' nhưng chưa có câu trả lời")

    for ma, n in Counter(t.get("ma_tc") for t in tcs).items():
        if n > 1:
            loi.append(f"Mã TC trùng: {ma}")
    tap_yc, tap_dx = set(ma_yc), set(ma_dx)
    for t in tcs:
        ma = t.get("ma_tc", "?")
        for truong in TRUONG_TC_BAT_BUOC:
            if not t.get(truong):
                loi.append(f"{ma}: thiếu '{truong}'")
        for y in t.get("ma_yeu_cau", []):
            if y not in tap_yc:
                loi.append(f"{ma}: mã yêu cầu '{y}' không có trong danh sách yêu cầu")
        for d in t.get("ma_dx", []):
            if d not in tap_dx:
                loi.append(f"{ma}: mã DX '{d}' không tồn tại")
            elif d in dx_da_chot:
                loi.append(f"{ma}: vẫn gắn {d} dù DX đã có trả lời – cập nhật TC theo trả lời rồi bỏ {d} khỏi ma_dx")
        if t.get("ma_dx") and "[CẦN XÁC NHẬN" not in t.get("ket_qua_mong_doi", ""):
            loi.append(f"{ma}: có DX liên quan nhưng kết quả mong đợi thiếu '[CẦN XÁC NHẬN'")
        if not t.get("ma_dx") and "[CẦN XÁC NHẬN" in t.get("ket_qua_mong_doi", ""):
            loi.append(f"{ma}: kết quả mong đợi còn '[CẦN XÁC NHẬN' nhưng ma_dx rỗng")
        if t.get("loai") not in LOAI_TC:
            loi.append(f"{ma}: loại không hợp lệ '{t.get('loai')}'")
        if t.get("ky_thuat") not in KY_THUAT:
            loi.append(f"{ma}: kỹ thuật không hợp lệ '{t.get('ky_thuat')}'")
        if t.get("uu_tien") not in UU_TIEN:
            loi.append(f"{ma}: ưu tiên không hợp lệ '{t.get('uu_tien')}'")
        if t.get("nhom") not in NHOM_TC:
            loi.append(f"{ma}: nhóm kịch bản không hợp lệ '{t.get('nhom')}' (Happy case / Chức năng / UI/UX / Hiệu năng)")
        if che_do == "Gọn" and t.get("uu_tien") == "Thấp":
            loi.append(f"{ma}: chế độ Gọn không được có TC ưu tiên Thấp")
    return loi


def ghi_bang(ws, tieu_de, dong, do_rong):
    ws.append(tieu_de)
    for o in ws[1]:
        o.font = Font(bold=True, color="FFFFFF")
        o.fill = MAU_TIEU_DE
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        o.border = VIEN
    for d in dong:
        ws.append(d)
    for hang in ws.iter_rows(min_row=2):
        for o in hang:
            o.alignment = Alignment(vertical="top", wrap_text=True)
            o.border = VIEN
    for i, w in enumerate(do_rong, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def xuat(du_lieu, duong_dan):
    tt = du_lieu.get("thong_tin", {})
    yeu_cau = du_lieu.get("yeu_cau", [])
    dx = du_lieu.get("diem_can_xac_nhan", [])
    tcs = du_lieu.get("test_cases", [])

    wb = Workbook()

    ws = wb.active
    ws.title = "Thông tin"
    nhan = [("Mã dự án", "ma_du_an"), ("Module", "ten_module"), ("Tài liệu nguồn", "tai_lieu_nguon"),
            ("Phiên bản tài liệu", "phien_ban_tai_lieu"), ("Người lập", "nguoi_lap"), ("Ngày lập", "ngay_lap")]
    for ten, khoa in nhan:
        ws.append([ten, tt.get(khoa, "")])
    ws.append(["Chế độ", tt.get("che_do", "Đầy đủ")])
    if tt.get("lan_cap_nhat"):
        ws.append(["Lần cập nhật", tt["lan_cap_nhat"]])
    ws.append([])
    ws.append(["Số yêu cầu", len(yeu_cau)])
    tt_dx = Counter(d.get("trang_thai", TT_DX_CHO) for d in dx)
    ws.append(["Số điểm cần xác nhận", f"{len(dx)} (" + ", ".join(f"{k}: {v}" for k, v in sorted(tt_dx.items())) + ")"])
    ws.append(["Số test case", len(tcs)])
    for loai, n in sorted(Counter(t["loai"] for t in tcs).items()):
        ws.append([f"  – {loai}", n])
    ws.append([])
    ws.append(["Lưu ý", "Sinh bởi AI. Tester phải kiểm tra lại các dòng có [CẦN XÁC NHẬN] trước khi dùng. "
                        "Điền trả lời của BA/KH vào sheet 'Điểm cần xác nhận' rồi chạy doc_tra_loi.py để cập nhật."])
    for hang in ws.iter_rows():
        hang[0].font = Font(bold=True)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 80

    ws = wb.create_sheet("Điểm cần xác nhận")
    ghi_bang(ws, [COT_DX_MA, "Vị trí", "Loại", "Nội dung", "Ảnh hưởng tới test", "Câu hỏi đề xuất",
                  "Mức ảnh hưởng", COT_DX_TRA_LOI, COT_DX_TRANG_THAI],
             [[d["ma"], d["vi_tri"], d["loai"], d["noi_dung"], d.get("anh_huong", ""), d["cau_hoi"],
               d["muc_anh_huong"], d.get("tra_loi", ""), d.get("trang_thai", TT_DX_CHO)] for d in dx],
             [8, 18, 14, 50, 40, 45, 12, 40, 14])
    for hang in ws.iter_rows(min_row=2):
        if hang[8].value == TT_DX_DA_CAP_NHAT:
            for o in hang:
                o.fill = MAU_DA_XONG
        elif hang[6].value == "Cao":
            for o in hang:
                o.fill = MAU_CANH_BAO

    ws = wb.create_sheet("Test case")
    ghi_bang(ws, [COT_TC_MA, "Mã yêu cầu", "Nhóm kịch bản", "Chức năng", "Tiêu đề", "Loại", "Kỹ thuật", "Ưu tiên",
                  "Điều kiện trước", "Bước thực hiện", "Dữ liệu test", "Kết quả mong đợi", "Mã DX",
                  *COT_TC_THUC_THI, "Cập nhật"],
             [[t["ma_tc"], ", ".join(t["ma_yeu_cau"]), t.get("nhom", ""), t["chuc_nang"], t["tieu_de"], t["loai"], t["ky_thuat"],
               t["uu_tien"], t["dieu_kien_truoc"],
               "\n".join(f"{i}. {b}" for i, b in enumerate(t["buoc_thuc_hien"], start=1)),
               t.get("du_lieu_test", ""), t["ket_qua_mong_doi"], ", ".join(t.get("ma_dx", [])),
               *(t.get(khoa) or ("Chưa chạy" if khoa == "trang_thai" else "") for khoa in COT_TC_THUC_THI.values()),
               t.get("cap_nhat", "")] for t in tcs],
             [14, 14, 13, 20, 40, 11, 9, 10, 30, 50, 30, 45, 10, 25, 11, 10, 18, 30])
    cot = {o.value: i for i, o in enumerate(ws[1])}
    for hang in ws.iter_rows(min_row=2):
        if hang[cot["Mã DX"]].value:
            hang[cot["Kết quả mong đợi"]].fill = MAU_CANH_BAO
        if hang[cot["Cập nhật"]].value:
            hang[cot["Cập nhật"]].fill = MAU_DA_XONG

    ws = wb.create_sheet("Ma trận truy vết")
    theo_yc = {y["ma"]: [] for y in yeu_cau}
    for t in tcs:
        for y in t["ma_yeu_cau"]:
            theo_yc.setdefault(y, []).append(t["ma_tc"])
    ghi_bang(ws, ["Mã yêu cầu", "Mô tả", "Nhóm", "Số TC", "Danh sách TC", "Ghi chú"],
             [[y["ma"], y["mo_ta"], y["nhom"], len(theo_yc[y["ma"]]), ", ".join(theo_yc[y["ma"]]),
               "CHƯA CÓ TEST CASE" if not theo_yc[y["ma"]] else ""] for y in yeu_cau],
             [12, 45, 18, 8, 70, 22])
    for hang in ws.iter_rows(min_row=2):
        if hang[3].value == 0:
            for o in hang:
                o.fill = MAU_THIEU

    wb.save(duong_dan)
    return [ma for ma, ds in theo_yc.items() if not ds]


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    sys.stdout.reconfigure(encoding="utf-8")
    with open(sys.argv[1], encoding="utf-8") as f:
        du_lieu = json.load(f)

    loi = kiem_tra(du_lieu)
    if loi:
        print(f"PHÁT HIỆN {len(loi)} LỖI – chưa xuất file:")
        for l in loi:
            print(f"  - {l}")
        sys.exit(1)

    chua_co_tc = xuat(du_lieu, sys.argv[2])
    tcs = du_lieu["test_cases"]
    dx = du_lieu["diem_can_xac_nhan"]
    print(f"Đã xuất: {sys.argv[2]}  (chế độ: {du_lieu.get('thong_tin', {}).get('che_do', 'Đầy đủ')})")
    print(f"  Yêu cầu: {len(du_lieu['yeu_cau'])} | Điểm cần xác nhận: {len(dx)} | Test case: {len(tcs)}")
    print("  Theo loại: " + ", ".join(f"{k} {v}" for k, v in sorted(Counter(t['loai'] for t in tcs).items())))
    print("  Theo ưu tiên: " + ", ".join(f"{k} {v}" for k, v in sorted(Counter(t['uu_tien'] for t in tcs).items())))
    print("  DX theo trạng thái: " + ", ".join(
        f"{k} {v}" for k, v in sorted(Counter(d.get('trang_thai', TT_DX_CHO) for d in dx).items())))
    print(f"  TC có [CẦN XÁC NHẬN]: {sum(1 for t in tcs if t.get('ma_dx'))}")
    if any(t.get("cap_nhat") for t in tcs):
        print(f"  TC có ghi chú cập nhật: {sum(1 for t in tcs if t.get('cap_nhat'))}")
    if chua_co_tc:
        print(f"  CẢNH BÁO – yêu cầu chưa có TC: {', '.join(chua_co_tc)}")
    else:
        print("  Mọi yêu cầu đều có ít nhất 1 TC.")


if __name__ == "__main__":
    main()
