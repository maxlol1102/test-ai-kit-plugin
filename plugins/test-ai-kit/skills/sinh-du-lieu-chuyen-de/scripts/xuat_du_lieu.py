"""Xuất dữ liệu kiểm thử ra Excel và tự kiểm tra.

Dùng: python xuat_du_lieu.py <du-lieu.json> <du-lieu-kiem-thu.xlsx>
Schema du-lieu.json: xem ../SKILL.md (mục "Định dạng du-lieu.json").
In ra JSON: {"so_bo", "so_dong", "loi": [...]} – còn lỗi thì sửa JSON rồi chạy lại.
"""
import json
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

LOAI = ("Hợp lệ", "Biên", "Sai", "Lớn")
CCCD_THAT = re.compile(r"^(?!000)\d{12}$")   # CCCD giả lập phải bắt đầu bằng 000
SDT_THAT = re.compile(r"^0(?!900000)\d{9}$")  # SĐT giả lập: 0900000xxx
# Tên cột có dấu hoặc không dấu, viết liền (SoCCCD, SoDienThoai, so_dien_thoai…)
COT_CCCD = re.compile(r"cccd|cmnd|c[aă]n.?c[uư][oớ]c|đ[iị]nh.?danh|dinh.?danh", re.I)
COT_SDT = re.compile(r"đi[eệ]n.?tho[aạ]i|dien.?thoai|sđt|sdt|phone|mobile|di.?d[oộ]ng", re.I)


def kiem_tra(dl):
    loi = []
    for k in ("chuyen_de", "quy_uoc", "bo_du_lieu", "buoc_thuc_hien", "yeu_cau_ky_thuat"):
        if k not in dl:
            loi.append(f"Thiếu khoá '{k}'")
    for i, bo in enumerate(dl.get("bo_du_lieu", []), 1):
        ten = bo.get("ten", f"bộ {i}")
        if bo.get("loai") not in LOAI:
            loi.append(f"{ten}: loại '{bo.get('loai')}' không hợp lệ ({', '.join(LOAI)})")
        cot = bo.get("cot", [])
        for j, d in enumerate(bo.get("dong", []), 1):
            if len(d) != len(cot):
                loi.append(f"{ten}: dòng {j} có {len(d)} ô, cần {len(cot)}")
            for c, o in zip(cot, d):
                s = str(o)
                if bo.get("loai") != "Sai" and COT_CCCD.search(c) and CCCD_THAT.match(s):
                    loi.append(f"{ten}: dòng {j} cột '{c}' = {s} giống CCCD thật – dùng số bắt đầu 000")
                if bo.get("loai") != "Sai" and COT_SDT.search(c) and SDT_THAT.match(s):
                    loi.append(f"{ten}: dòng {j} cột '{c}' = {s} giống SĐT thật – dùng 0900000xxx")
        if len(bo.get("dong", [])) > 50:
            loi.append(f"{ten}: quá 50 dòng mẫu – dùng 'so_dong_sinh_them' cho dữ liệu lớn")
    return loi


def xuat(dl, duong_dan):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng quan"
    for d in (["DỮ LIỆU KIỂM THỬ"], ["Chuyên đề", dl.get("chuyen_de", "")], ["Quy ước giả lập", dl.get("quy_uoc", "")],
              ["Cần xác nhận", "\n".join(f"- {x}" for x in dl.get("can_xac_nhan", []))], []):
        ws.append(d)
    ws.append(["Sheet", "Bộ dữ liệu", "Loại", "Bảng liên quan", "Số dòng", "TC liên quan", "Mô tả"])
    hang_tieu_de = ws.max_row
    for i, bo in enumerate(dl["bo_du_lieu"], 1):
        ten_sheet = f"{i:02d}-{bo['loai']}"[:31]
        ws.append([ten_sheet, bo["ten"], bo["loai"], bo.get("bang_lien_quan", ""),
                   len(bo["dong"]) + bo.get("so_dong_sinh_them", 0), ", ".join(bo.get("tc_lien_quan", [])), bo.get("mo_ta", "")])
        s = wb.create_sheet(ten_sheet)
        s.append(bo["cot"])
        for d in bo["dong"]:
            s.append(d)
        mau = bo["dong"] or [[""] * len(bo["cot"])]
        for j in range(bo.get("so_dong_sinh_them", 0)):  # dữ liệu lớn: nhân bản dòng mẫu, đánh số cột đầu để không trùng
            dong = list(mau[j % len(mau)])
            if dong:
                dong[0] = f"{dong[0]}-{j + 1:05d}"
            s.append(dong)
        for o in s[1]:
            o.font = Font(bold=True)
        for k in range(1, s.max_column + 1):
            s.column_dimensions[s.cell(1, k).column_letter].width = 22
    for ten_sheet, khoa, tieu_de in (("Bước thực hiện", "buoc_thuc_hien", "Bước"), ("Yêu cầu kỹ thuật", "yeu_cau_ky_thuat", "Yêu cầu")):
        s = wb.create_sheet(ten_sheet)
        s.append(["STT", tieu_de])
        for j, x in enumerate(dl.get(khoa, []), 1):
            s.append([j, x])
        s.column_dimensions["B"].width = 100
        for o in s[1]:
            o.font = Font(bold=True)
    if dl.get("van_ban"):
        s = wb.create_sheet("Văn bản mẫu")
        s.append(["File", "Loại văn bản", "Dùng cho", "TC liên quan"])
        for v in dl["van_ban"]:
            s.append([v.get("file", ""), v.get("loai", ""), v.get("dung_cho", ""), ", ".join(v.get("tc_lien_quan", []))])
        for o in s[1]:
            o.font = Font(bold=True)
    for o in ws[hang_tieu_de]:
        o.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=13)
    for hang in ws.iter_rows():
        for o in hang:
            o.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["G"].width = 50
    wb.save(duong_dan)


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    dl = json.load(open(sys.argv[1], encoding="utf-8"))
    loi = kiem_tra(dl)
    if not any(l.startswith("Thiếu khoá 'bo_du_lieu'") for l in loi):
        xuat(dl, sys.argv[2])
    print(json.dumps({"so_bo": len(dl.get("bo_du_lieu", [])),
                      "so_dong": sum(len(b.get("dong", [])) + b.get("so_dong_sinh_them", 0) for b in dl.get("bo_du_lieu", [])),
                      "loi": loi}, ensure_ascii=False))
    sys.exit(1 if loi else 0)


if __name__ == "__main__":
    main()
