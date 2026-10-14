"""Tách tài liệu Word lớn theo cấu trúc chuẩn của Trí Nam để agent xử lý từng phần.

SRS (Đặc tả yêu cầu người sử dụng):
    Phần chung (I–III) + Yêu cầu phi chức năng (V)
    IV. YÊU CẦU CHI TIẾT VỀ CHỨC NĂNG
        IV.x  PHÂN HỆ ...                       ← đoạn bắt đầu bằng "IV.<số>"
            N. <Tên chức năng>                  ← đoạn đứng ngay trước "a. Thông tin chức năng ..."
                a. bảng Tên chức năng / Tác nhân / Luồng / Quy tắc nghiệp vụ
                b. Giao diện (ảnh) + bảng phần tử giao diện

CSDL: các bảng trường dữ liệu → chữ gọn (gộp ô trộn).
"""
import io
import json
import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image

NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
CANH_ANH_TOI_DA = 1568
ANH_TOI_DA_MOI_CN = 4


def _dong_bang(hang):
    o = []
    for c in hang.cells:
        x = re.sub(r"\s*\n\s*", " / ", c.text.strip())
        if not o or o[-1] != x:
            o.append(x)
    return " | ".join(o)


def _cac_khoi(tai_lieu):
    """Duyệt thân tài liệu đúng thứ tự: ("p", chữ, [rId ảnh], None) hoặc ("t", chữ của bảng, [rId ảnh], Table)."""
    for el in tai_lieu.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, tai_lieu)
            anh = [b.get(f"{NS_R}embed") for b in el.iter(f"{NS_A}blip") if b.get(f"{NS_R}embed")]
            chu = p.text.strip()
            if chu or anh:
                yield "p", chu, anh, None
        elif tag == "tbl":
            t = Table(el, tai_lieu)
            anh = [b.get(f"{NS_R}embed") for b in el.iter(f"{NS_A}blip") if b.get(f"{NS_R}embed")]
            yield "t", "\n".join(_dong_bang(h) for h in t.rows), anh, t


# ---------------------------------------------------------------- đọc bảng của một chức năng thành dữ liệu có cấu trúc
# Đọc thẳng từ ô Word (giữ đúng cột) – chuỗi van_ban gộp ô trùng nên lệch cột, không dùng để tách trường được.

# Loại phần tử mà người dùng nhập/chọn giá trị (đầu vào của chức năng). Button, Label, Pagination… không phải đầu vào.
LOAI_NHAP = {"text field", "textarea", "text area", "search field", "dropdown list", "dropdown", "combobox",
             "date picker", "checkbox", "check box", "chekcbox", "radio button", "radio", "toggles", "toggle",
             "upload file", "file upload"}
_TRUONG_THONG_TIN = {"Mô tả chức năng": "mo_ta", "Điều kiện trước": "dieu_kien_truoc", "Điều kiện sau": "dieu_kien_sau"}


def _o(hang):
    return [c.text.strip() for c in hang.cells]


def _doc_bang_thong_tin(bang):
    """Bảng "a. Thông tin chức năng và luồng sự kiện" (4 cột, ô trộn lặp lại giá trị)."""
    tt = {"mo_ta": "", "tac_nhan": [], "dieu_kien_truoc": "", "dieu_kien_sau": "",
          "luong_chinh": [], "luong_ngoai_le": [], "quy_tac": []}
    ngoai_le = False
    for hang in bang.rows:
        o = _o(hang)
        dau, gia_tri = o[0], (o[1] if len(o) > 1 else "")
        if dau in _TRUONG_THONG_TIN:
            tt[_TRUONG_THONG_TIN[dau]] = gia_tri
        elif dau == "Tác nhân":
            tt["tac_nhan"] = [x.strip() for x in re.split(r"[,;]", gia_tri) if x.strip()]
        elif dau == "Luồng sự kiện ngoại lệ":
            ngoai_le = True
            if gia_tri and gia_tri != dau:
                tt["luong_ngoai_le"].append(gia_tri)
        elif dau == "Quy tắc nghiệp vụ":
            ngoai_le = False
            m = re.match(r"^(BR\d+)\s*:?\s*(.*)$", gia_tri, re.S)
            tt["quy_tac"].append({"ma": m.group(1), "noi_dung": m.group(2).strip()} if m else {"ma": "", "noi_dung": gia_tri})
        elif re.match(r"^\(\d+\)", dau):
            if ngoai_le:
                tt["luong_ngoai_le"].append(" – ".join(dict.fromkeys(x for x in o if x)))
            else:
                tt["luong_chinh"].append({"buoc": dau, "he_thong": o[2] if len(o) > 2 else "",
                                          "thay_the": o[3] if len(o) > 3 and o[3] != o[2] else ""})
    return tt


def _doc_bang_giao_dien(bang):
    """Bảng "Mô tả giao diện người dùng (GUI elements)" → [{ten màn, phan_tu: [...]}].
    Dòng trộn hết cột = tên màn hình; dòng chỉ có ô đầu = nhóm thông tin trong màn; dòng mẫu trống (chỉ có số TT) bỏ qua."""
    dau = _o(bang.rows[0])

    def cot(tien_to):
        return next((i for i, x in enumerate(dau) if x.startswith(tien_to)), None)

    vt = {"ten": cot("Tên"), "loai": cot("Loại phần tử"), "mac_dinh": cot("Giá trị mặc định"),
          "bat_buoc": cot("Bắt buộc"), "mo_ta": cot("Mô tả")}
    man, man_hinh, nhom = None, [], ""
    for hang in bang.rows[1:]:
        o = _o(hang)
        if not any(o):
            continue
        if len(set(o)) == 1:
            man, nhom = {"ten": o[0], "phan_tu": []}, ""
            man_hinh.append(man)
            continue
        pt = {k: (o[i] if i is not None and i < len(o) else "") for k, i in vt.items()}
        if not pt["loai"] and not pt["ten"]:
            if o[0] and not o[0].isdigit():
                nhom = o[0]
            continue
        if not pt["loai"] and not any(o[1:]):
            nhom = o[0]
            continue
        if man is None:
            man = {"ten": "", "phan_tu": []}
            man_hinh.append(man)
        pt["nhom"] = nhom
        pt["la_dau_vao"] = pt["loai"].lower() in LOAI_NHAP
        man["phan_tu"].append(pt)
    return [m for m in man_hinh if m["phan_tu"]]


def _cau_truc_chuc_nang(bang_cua_cn):
    """Gom bảng thông tin + các bảng giao diện của một chức năng; liệt kê chỗ tài liệu còn thiếu."""
    tt, man_hinh = None, []
    for bang in bang_cua_cn:
        dau = _o(bang.rows[0])
        if dau and dau[0] == "Tên chức năng":
            tt = _doc_bang_thong_tin(bang)
        elif any(x.startswith("Loại phần tử") for x in dau):
            man_hinh.extend(_doc_bang_giao_dien(bang))
    thieu = []
    if tt is None:
        thieu.append("Không có bảng Thông tin chức năng")
        tt = {"mo_ta": "", "tac_nhan": [], "dieu_kien_truoc": "", "dieu_kien_sau": "",
              "luong_chinh": [], "luong_ngoai_le": [], "quy_tac": []}
    else:
        if not tt["luong_chinh"]:
            thieu.append("Luồng sự kiện chính trống")
        if not tt["luong_ngoai_le"]:
            thieu.append("Luồng sự kiện ngoại lệ trống")
        br_trong = [q["ma"] or "BR?" for q in tt["quy_tac"] if not q["noi_dung"]]
        if br_trong:
            thieu.append("Quy tắc nghiệp vụ để trống: " + ", ".join(br_trong))
    dau_vao = [pt for m in man_hinh for pt in m["phan_tu"] if pt["la_dau_vao"]]
    if not man_hinh:
        thieu.append("Chưa mô tả phần tử giao diện (bảng GUI trống hoặc không có)")
    elif not dau_vao:
        thieu.append("Không có trường nhập liệu nào trong bảng GUI")
    chua_ro = [pt["ten"] or pt["loai"] for pt in dau_vao if pt["bat_buoc"] not in ("Y", "N")]
    if chua_ro:
        thieu.append(f"{len(chua_ro)}/{len(dau_vao)} trường nhập chưa ghi Bắt buộc Y/N")
    return {**tt, "man_hinh": man_hinh}, thieu


def _luu_anh(tai_lieu, rid, duong_dan):
    """Lưu ảnh PNG đã thu nhỏ. Bỏ qua định dạng không đọc được (EMF/WMF...)."""
    try:
        phan = tai_lieu.part.related_parts[rid]
        anh = Image.open(io.BytesIO(phan.blob))
        anh.load()
    except Exception:  # noqa: BLE001 – ảnh lỗi/không hỗ trợ thì bỏ qua
        return False
    if anh.width < 200 or anh.height < 120:
        return False  # biểu tượng nhỏ, không phải màn hình
    anh.thumbnail((CANH_ANH_TOI_DA, CANH_ANH_TOI_DA))
    anh.convert("RGB").save(duong_dan, "PNG", optimize=True)
    return True


def tach_srs(duong_dan_docx, thu_muc_ra):
    """Trả về cấu trúc {chung, phi_chuc_nang, phan_he:[{ma, ten, chuc_nang:[...]}]} và lưu ảnh giao diện."""
    tl = Document(duong_dan_docx)
    khoi = list(_cac_khoi(tl))
    chu_p = [k[1] if k[0] == "p" else None for k in khoi]
    # Bỏ qua mục lục: phần chi tiết bắt đầu ở lần xuất hiện CUỐI của tiêu đề (bản trong mục lục có số trang dính kèm)
    vt_chi_tiet = max((i for i, c in enumerate(chu_p) if c and c.upper() == "YÊU CẦU CHI TIẾT VỀ CHỨC NĂNG"), default=None)
    vt_phi_cn = max((i for i, c in enumerate(chu_p) if c and c.upper() == "YÊU CẦU PHI CHỨC NĂNG"), default=len(khoi))
    if vt_chi_tiet is None:
        raise ValueError("Không thấy mục 'YÊU CẦU CHI TIẾT VỀ CHỨC NĂNG' – tài liệu không theo mẫu SRS chuẩn.")

    def van_ban(ds):
        return "\n".join(k[1] for k in ds if k[1])

    thu_muc_anh = Path(thu_muc_ra) / "anh"
    thu_muc_anh.mkdir(parents=True, exist_ok=True)
    phan_he, ph, cn = [], None, None
    for i in range(vt_chi_tiet + 1, vt_phi_cn):
        loai, chu, anh, _ = khoi[i]
        m_ph = re.match(r"^(IV\.\d+(?:\.\d+)*)\.?\s+(.+)$", chu) if loai == "p" else None
        if m_ph:
            ph = {"ma": m_ph.group(1), "ten": m_ph.group(2).strip(), "chuc_nang": []}
            phan_he.append(ph)
            cn = None
            continue
        ke_tiep = next((chu_p[j] for j in range(i + 1, min(i + 4, vt_phi_cn)) if chu_p[j]), "") or ""
        m_cn = re.match(r"^(\d+)\.\s*(.+)$", chu) if loai == "p" else None
        if m_cn and ke_tiep.lower().startswith("a. thông tin chức năng"):
            if ph is None:
                ph = {"ma": "IV", "ten": "Chức năng", "chuc_nang": []}
                phan_he.append(ph)
            stt = int(m_cn.group(1))
            cn = {"ma": f"{ph['ma']}-CN{stt:02d}", "stt": stt, "ten": m_cn.group(2).strip(), "khoi": [],
                  "rid_gd": [], "rid_bang": []}
            ph["chuc_nang"].append(cn)
            continue
        if cn is not None:
            cn["khoi"].append(khoi[i])
            # Ảnh màn hình (đoạn văn ở mục Giao diện) ưu tiên hơn ảnh trong bảng (thường là biểu đồ luồng)
            cn["rid_gd" if loai == "p" else "rid_bang"].extend(anh)
    for p in phan_he:
        for c in p["chuc_nang"]:
            c["anh"] = []
            for rid in dict.fromkeys(c.pop("rid_gd") + c.pop("rid_bang")):
                if len(c["anh"]) >= ANH_TOI_DA_MOI_CN:
                    break
                duong_dan = thu_muc_anh / f"{c['ma']}-{len(c['anh']) + 1}.png"
                if _luu_anh(tl, rid, duong_dan):
                    c["anh"].append(duong_dan.name)
            c["thong_tin"], c["thieu"] = _cau_truc_chuc_nang([k[3] for k in c["khoi"] if k[0] == "t"])
            c["van_ban"] = f"{c['stt']}. {c['ten']}\n" + van_ban(c.pop("khoi"))
            c["so_ky_tu"] = len(c["van_ban"])
    ket_qua = {
        "chung": van_ban(khoi[:vt_chi_tiet]),
        "phi_chuc_nang": van_ban(khoi[vt_phi_cn:]),
        "phan_he": [p for p in phan_he if p["chuc_nang"]],
    }
    ket_qua["tong_chuc_nang"] = sum(len(p["chuc_nang"]) for p in ket_qua["phan_he"])
    return ket_qua


def tach_csdl(duong_dan_docx):
    """Bảng thiết kế CSDL → chữ gọn, giữ thứ tự (tên bảng ở đoạn ngay trước mỗi bảng trường)."""
    tl = Document(duong_dan_docx)
    return "\n".join(chu if loai == "p" else f"[Bảng]\n{chu}" for loai, chu, _, _ in _cac_khoi(tl) if chu)


def la_srs_chuan(duong_dan_docx):
    try:
        tl = Document(duong_dan_docx)
    except Exception:  # noqa: BLE001
        return False
    return any(p.text.strip().upper() == "YÊU CẦU CHI TIẾT VỀ CHỨC NĂNG" for p in tl.paragraphs)


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    kq = tach_srs(sys.argv[1], sys.argv[2])
    print(json.dumps({"chung": len(kq["chung"]), "phi_chuc_nang": len(kq["phi_chuc_nang"]), "tong": kq["tong_chuc_nang"]}))
