"""Thống kê kết quả kiểm thử (bằng code – AI không tự đếm) và xuất Biểu mẫu kết quả theo Mẫu 3 Trí Nam.

Dùng:
  python xuat_bao_cao.py thong-ke <thư mục dự án>                  → in JSON số liệu (để agent viết đánh giá)
  python xuat_bao_cao.py excel    <thư mục dự án> <file .xlsx ra>   → biểu mẫu Excel (đọc bao-cao/danh-gia.json nếu có)
Đọc: trang-thai.json, kich-ban/test-case.json, bao-cao/loi.json, bao-cao/danh-gia.json (trong thư mục dự án).
"""
import json
import sys
from collections import Counter
from datetime import date
from pathlib import Path

TT_TC = ("Đạt", "Không đạt", "Bị chặn", "Chưa chạy")
TT_LOI = ("Mới", "Đang sửa", "Đã sửa", "Đóng", "Mở lại", "Từ chối")
MUC_DO = ("Critical", "Major", "Minor", "Trivial")


def _doc(p, mac_dinh):
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else mac_dinh


def thong_ke(da):
    da = Path(da)
    tcs = _doc(da / "kich-ban" / "test-case.json", {}).get("test_cases", [])
    loi = _doc(da / "bao-cao" / "loi.json", [])
    tt = lambda t: t.get("trang_thai") or "Chưa chạy"  # noqa: E731
    theo_nhom = {}
    for t in tcs:
        theo_nhom.setdefault(t.get("nhom", "Khác"), Counter())[tt(t)] += 1
    theo_cn = {}
    for t in tcs:
        theo_cn.setdefault(t.get("chuc_nang", ""), Counter())[tt(t)] += 1
    cao = [t for t in tcs if t.get("uu_tien") == "Cao"]
    return {
        "tong_tc": len(tcs),
        "theo_trang_thai": dict(Counter(tt(t) for t in tcs)),
        "ty_le_dat": round(sum(1 for t in tcs if tt(t) == "Đạt") / len(tcs), 3) if tcs else 0,
        "tc_cao_chua_chay": sum(1 for t in cao if tt(t) == "Chưa chạy"), "tong_tc_cao": len(cao),
        "ket_qua_ai_chua_xac_nhan": sum(1 for t in tcs if str(t.get("nguon_ket_qua", "")).startswith("AI")),
        "theo_nhom": {k: dict(v) for k, v in theo_nhom.items()},
        "theo_chuc_nang": {k: dict(v) for k, v in theo_cn.items()},
        "tong_loi": len(loi),
        "loi_theo_muc_do": dict(Counter(l.get("muc_do") for l in loi)),
        "loi_nghiem_trong_mo": [l["ma_loi"] for l in loi if l.get("muc_do") in ("Critical", "Major")
                                and l.get("trang_thai") not in ("Đóng", "Từ chối")],
        "loi_chua_xac_nhan": sum(1 for l in loi if not l.get("da_xac_nhan")),
    }


def excel(da, ra):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    da = Path(da)
    st = _doc(da / "trang-thai.json", {})
    tcs = _doc(da / "kich-ban" / "test-case.json", {}).get("test_cases", [])
    loi = _doc(da / "bao-cao" / "loi.json", [])
    dg = _doc(da / "bao-cao" / "danh-gia.json", {})
    tk = thong_ke(da)
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    dong = [["BIỂU MẪU KẾT QUẢ KIỂM THỬ"], [],
            ["Dự án", f"{st.get('ma', '')} – {st.get('ten', '')}"], ["Môi trường", st.get("url", "")],
            ["Ngày lập", date.today().strftime("%d/%m/%Y")], ["Người lập", st.get("nguoi_phu_trach", "") + " (AI hỗ trợ – tester xác nhận)"],
            ["Phạm vi", dg.get("pham_vi", "")], [],
            ["KẾT QUẢ TEST CASE"], ["Nhóm kịch bản", "Tổng", *TT_TC, "Tỷ lệ đạt"]]
    for ten, c in tk["theo_nhom"].items():
        tong = sum(c.values())
        dong.append([ten, tong, *(c.get(k, 0) for k in TT_TC), f"{c.get('Đạt', 0) / tong:.0%}" if tong else ""])
    dong += [[], ["THỐNG KÊ LỖI"], ["Mức độ", *TT_LOI, "Tổng"]]
    for md in MUC_DO:
        c = Counter(l.get("trang_thai") for l in loi if l.get("muc_do") == md)
        dong.append([md, *(c.get(k, 0) for k in TT_LOI), sum(c.values())])
    dong += [[], ["ĐÁNH GIÁ"], ["Kết luận", dg.get("ket_luan", "(chưa có đánh giá)")], ["Đánh giá chung", dg.get("danh_gia_chung", "")],
             ["Lỗi tồn đọng nghiêm trọng", ", ".join(tk["loi_nghiem_trong_mo"]) or "Không"],
             ["Rủi ro", "\n".join(f"- {r}" for r in dg.get("rui_ro", []))], ["Kiến nghị", "\n".join(f"- {r}" for r in dg.get("kien_nghi", []))],
             ["Phần AI thực hiện", f"{tk['ket_qua_ai_chua_xac_nhan']} kết quả AI chưa được tester xác nhận; {tk['loi_chua_xac_nhan']} lỗi chưa xác nhận"],
             [], ["Lưu ý", "Kết quả do AI chạy tự động phải được tester xác nhận trước khi đưa vào hồ sơ nghiệm thu."]]
    for d in dong:
        ws.append(d)
    for hang in ws.iter_rows():
        for o in hang:
            o.alignment = Alignment(wrap_text=True, vertical="top")
        if hang[0].value and str(hang[0].value).isupper():
            hang[0].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70
    s = wb.create_sheet("Kết quả test case")
    s.append(["Mã TC", "Nhóm", "Chức năng", "Tiêu đề", "Ưu tiên", "Kết quả mong đợi", "Kết quả thực tế", "Trạng thái",
              "Mã lỗi", "Người thực hiện / Ngày", "Nguồn kết quả", "Bằng chứng"])
    for t in tcs:
        s.append([t["ma_tc"], t.get("nhom", ""), t.get("chuc_nang", ""), t.get("tieu_de", ""), t.get("uu_tien", ""),
                  t.get("ket_qua_mong_doi", ""), t.get("ket_qua_thuc_te", ""), t.get("trang_thai") or "Chưa chạy", t.get("ma_loi", ""),
                  t.get("nguoi_thuc_hien", ""), t.get("nguon_ket_qua", "Tester"), "\n".join(t.get("bang_chung", []))])
    s = wb.create_sheet("Danh sách lỗi")
    s.append(["Mã lỗi", "Tiêu đề", "Mức độ", "Ưu tiên", "Trạng thái", "Test case", "Bước tái hiện", "Kết quả mong đợi",
              "Kết quả thực tế", "Môi trường", "Người phát hiện", "Ngày", "Nguồn", "Đã xác nhận", "Bằng chứng"])
    for l in loi:
        s.append([l["ma_loi"], l.get("tieu_de", ""), l.get("muc_do", ""), l.get("uu_tien", ""), l.get("trang_thai", ""), l.get("ma_tc", ""),
                  "\n".join(f"{i}. {b}" for i, b in enumerate(l.get("buoc_tai_hien", []), 1)), l.get("ket_qua_mong_doi", ""),
                  l.get("ket_qua_thuc_te", ""), l.get("moi_truong", ""), l.get("nguoi_phat_hien", ""), l.get("ngay_phat_hien", ""),
                  l.get("nguon_phat_hien", ""), "Có" if l.get("da_xac_nhan") else "Chưa", "\n".join(l.get("bang_chung", []))])
    for sh in wb.worksheets[1:]:
        for o in sh[1]:
            o.font = Font(bold=True)
        for hang in sh.iter_rows(min_row=2):
            for o in hang:
                o.alignment = Alignment(wrap_text=True, vertical="top")
        for i in range(1, sh.max_column + 1):
            sh.column_dimensions[sh.cell(1, i).column_letter].width = 22
    Path(ra).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ra)
    return tk


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 3 or sys.argv[1] not in ("thong-ke", "excel") or (sys.argv[1] == "excel" and len(sys.argv) != 4):
        print(__doc__)
        sys.exit(2)
    kq = thong_ke(sys.argv[2]) if sys.argv[1] == "thong-ke" else excel(sys.argv[2], sys.argv[3])
    print(json.dumps(kq, ensure_ascii=False))


if __name__ == "__main__":
    main()
