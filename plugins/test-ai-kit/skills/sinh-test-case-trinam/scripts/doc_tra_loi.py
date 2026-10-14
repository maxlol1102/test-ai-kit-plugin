"""Đọc trả lời của BA/KH và kết quả chạy test từ file Excel, gộp vào JSON.

Cách dùng:
    python doc_tra_loi.py <file.xlsx> <json_cu.json> <json_moi.json>

- Sheet "Điểm cần xác nhận": cột "Trả lời của BA/KH" → tra_loi. DX đang "Chờ trả lời" mà có trả lời
  → chuyển "Đã trả lời". DX đã "Đã cập nhật TC" mà câu trả lời bị sửa → quay lại "Đã trả lời".
- Sheet "Test case": các cột tester điền khi chạy (Kết quả thực tế, Trạng thái, Mã lỗi, Người thực hiện)
  được chép vào JSON để lần xuất Excel sau không bị mất.
- In ra danh sách DX cần xử lý và các TC bị ảnh hưởng.

Script KHÔNG sửa nội dung test case – việc đó do agent làm theo references/cap-nhat-theo-tra-loi.md.
"""
import json
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xuat_excel import (COT_DX_MA, COT_DX_TRA_LOI, COT_TC_MA, COT_TC_THUC_THI, TT_DX_CHO,  # noqa: E402
                        TT_DX_DA_CAP_NHAT, TT_DX_DA_TRA_LOI)


def doc_sheet(wb, ten):
    if ten not in wb.sheetnames:
        sys.exit(f"Không thấy sheet '{ten}' trong file Excel.")
    hang = list(wb[ten].iter_rows(values_only=True))
    tieu_de = [str(o).strip() if o is not None else "" for o in hang[0]]
    return [dict(zip(tieu_de, h)) for h in hang[1:] if any(o not in (None, "") for o in h)]


def chuoi(o):
    return "" if o is None else str(o).strip()


def main():
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(2)
    sys.stdout.reconfigure(encoding="utf-8")
    duong_xlsx, duong_cu, duong_moi = sys.argv[1:]
    wb = load_workbook(duong_xlsx, read_only=True, data_only=True)
    with open(duong_cu, encoding="utf-8") as f:
        du_lieu = json.load(f)

    # 1. Trả lời DX
    dx_theo_ma = {d["ma"]: d for d in du_lieu["diem_can_xac_nhan"]}
    can_xu_ly = []
    for dong in doc_sheet(wb, "Điểm cần xác nhận"):
        ma, tra_loi = chuoi(dong.get(COT_DX_MA)), chuoi(dong.get(COT_DX_TRA_LOI))
        d = dx_theo_ma.get(ma)
        if d is None:
            print(f"CẢNH BÁO – DX '{ma}' có trong Excel nhưng không có trong JSON, bỏ qua.")
            continue
        if tra_loi == chuoi(d.get("tra_loi")):
            continue
        d["tra_loi"] = tra_loi
        if not tra_loi:
            d["trang_thai"] = TT_DX_CHO
        elif d.get("trang_thai", TT_DX_CHO) in (TT_DX_CHO, TT_DX_DA_CAP_NHAT):
            d["trang_thai"] = TT_DX_DA_TRA_LOI
        if tra_loi:
            can_xu_ly.append(d)

    # 2. Kết quả chạy test
    tc_theo_ma = {t["ma_tc"]: t for t in du_lieu["test_cases"]}
    so_tc_co_ket_qua = 0
    for dong in doc_sheet(wb, "Test case"):
        ma = chuoi(dong.get(COT_TC_MA))
        t = tc_theo_ma.get(ma)
        if t is None:
            print(f"CẢNH BÁO – TC '{ma}' có trong Excel nhưng không có trong JSON (tester tự thêm?), cần xử lý tay.")
            continue
        co_ket_qua = False
        for cot, khoa in COT_TC_THUC_THI.items():
            gia_tri = chuoi(dong.get(cot))
            if khoa == "trang_thai" and gia_tri == "Chưa chạy":
                gia_tri = ""
            if gia_tri:
                t[khoa] = gia_tri
                co_ket_qua = True
            else:
                t.pop(khoa, None)
        so_tc_co_ket_qua += co_ket_qua

    with open(duong_moi, "w", encoding="utf-8") as f:
        json.dump(du_lieu, f, ensure_ascii=False, indent=2)

    print(f"Đã ghi: {duong_moi}")
    print(f"  TC có kết quả chạy được giữ lại: {so_tc_co_ket_qua}")
    print(f"  DX có trả lời mới: {len(can_xu_ly)}")
    for d in can_xu_ly:
        tc_lien_quan = [t["ma_tc"] for t in du_lieu["test_cases"] if d["ma"] in t.get("ma_dx", [])]
        print(f"\n  [{d['ma']}] {d['vi_tri']} – {d['noi_dung']}")
        print(f"    Trả lời: {d['tra_loi']}")
        print(f"    TC đang gắn {d['ma']} ({len(tc_lien_quan)}): {', '.join(tc_lien_quan) or '(không có)'}")
    con_cho = [d["ma"] for d in du_lieu["diem_can_xac_nhan"] if d.get("trang_thai", TT_DX_CHO) == TT_DX_CHO]
    print(f"\n  DX còn chờ trả lời ({len(con_cho)}): {', '.join(con_cho)}")


if __name__ == "__main__":
    main()
