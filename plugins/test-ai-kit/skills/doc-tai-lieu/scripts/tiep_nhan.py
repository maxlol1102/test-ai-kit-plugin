"""Tiếp nhận đầu vào tester gửi: chép file/thư mục vào <dự án>/dau-vao/, phân loại từng file, ghi dau-vao/danh-muc.json.

Dùng: python tiep_nhan.py <thư mục dự án> [<file hoặc thư mục> ...]
  Không truyền đường dẫn → chỉ quét lại những gì đã có trong dau-vao/.
In JSON: {"so_file", "theo_loai": {...}, "file": [{"file", "loai", "can_cu", "canh_bao"}], "thieu": [...]}
Loại: SRS | CSDL | Yêu cầu kiểm thử | Dữ liệu mẫu | Checklist thiết lập | Ảnh giao diện | Khác
"""
import json
import re
import shutil
import sys
import unicodedata
from pathlib import Path

ANH = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
VAN_BAN = {".docx", ".pdf", ".md", ".txt", ".doc"}
BANG = {".xlsx", ".xls", ".csv"}
BO_QUA = {"danh-muc.json", "yeu-cau.md", "desktop.ini", "thumbs.db", ".ds_store"}
TEN = [  # (loại, mẫu trên tên file không dấu) – xét theo thứ tự
    ("Checklist thiết lập", r"checklist|thiet.?lap|cau.?hinh|setup|tai.?khoan|moi.?truong"),
    ("Yêu cầu kiểm thử", r"yeu.?cau.?kiem.?thu|ke.?hoach.?kiem.?thu|test.?plan|pham.?vi|de.?nghi.?kiem.?thu|yeu.?cau.?test"),
    ("CSDL", r"csdl|co.?so.?du.?lieu|database|\bdb\b|erd|thiet.?ke.?du.?lieu|mo.?hinh.?du.?lieu"),
    ("SRS", r"srs|dac.?ta|yeu.?cau.?phan.?mem|\burd\b|use.?case|tai.?lieu.?yeu.?cau|phan.?tich.?thiet.?ke|hdsd|huong.?dan.?su.?dung"),
    ("Dữ liệu mẫu", r"du.?lieu|data|mau|danh.?muc|import"),
]
COT_CCCD = re.compile(r"cccd|cmnd|can.?cuoc|dinh.?danh", re.I)
COT_SDT = re.compile(r"dien.?thoai|sdt|phone|mobile|di.?dong", re.I)


def khong_dau(s):
    s = unicodedata.normalize("NFD", s.replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()


def _chu_dau(p, gioi_han=6000):
    """Vài nghìn ký tự đầu của văn bản để đoán loại."""
    try:
        if p.suffix.lower() == ".docx":
            import docx
            d = docx.Document(str(p))
            chu = "\n".join(x.text for x in d.paragraphs[:400])
            so_bang_kieu = sum(1 for t in d.tables[:80] if t.rows and "kieu" in khong_dau(" ".join(c.text for c in t.rows[0].cells)))
            return chu[:gioi_han], so_bang_kieu
        if p.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(str(p))
            return "\n".join((pg.extract_text() or "") for pg in r.pages[:6])[:gioi_han], 0
        if p.suffix.lower() in (".md", ".txt"):
            return p.read_text(encoding="utf-8", errors="ignore")[:gioi_han], 0
    except Exception:  # file hỏng/khoá mật khẩu → đoán theo tên
        pass
    return "", 0


def _du_lieu_that(p):
    """Bảng có cột CCCD/SĐT mà giá trị trông như thật → cảnh báo dữ liệu cá nhân."""
    try:
        dong = []
        if p.suffix.lower() == ".csv":
            import csv
            with open(p, encoding="utf-8-sig", errors="ignore") as f:
                dong = [r for _, r in zip(range(30), csv.reader(f))]
        elif p.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(p), read_only=True, data_only=True)
            for ws in wb.worksheets[:5]:
                dong += [[("" if v is None else str(v)) for v in r] for r in ws.iter_rows(max_row=30, values_only=True)]
        if not dong:
            return False
        tieu_de = [khong_dau(str(c)) for c in dong[0]]
        for i, c in enumerate(tieu_de):
            if COT_CCCD.search(c) or COT_SDT.search(c):
                gia_tri = [str(r[i]).strip() for r in dong[1:] if i < len(r)]
                if any(re.fullmatch(r"(?!000)\d{12}", g) or re.fullmatch(r"0(?!900000)\d{9}", g) for g in gia_tri):
                    return True
    except Exception:
        pass
    return False


def phan_loai(p):
    ten = khong_dau(p.stem)
    duoi = p.suffix.lower()
    canh_bao = []
    if duoi in ANH:
        return "Ảnh giao diện", "định dạng ảnh", canh_bao
    if duoi == ".doc" or duoi == ".xls":
        canh_bao.append("Định dạng cũ – mở bằng Word/Excel và lưu lại thành .docx/.xlsx để đọc được")
    loai, can_cu = None, ""
    for l, mau in TEN:
        if re.search(mau, ten) and not (l == "Dữ liệu mẫu" and duoi not in BANG):
            loai, can_cu = l, f"tên file khớp '{l}'"
            break
    if duoi in VAN_BAN and loai in (None, "SRS", "CSDL"):
        chu, so_bang_kieu = _chu_dau(p)
        k = khong_dau(chu)
        if "yeu cau chi tiet ve chuc nang" in k or ("use case" in k and "tac nhan" in k):
            loai, can_cu = "SRS", "nội dung có mục yêu cầu chức năng / use case"
        elif so_bang_kieu >= 3 or "thiet ke co so du lieu" in k:
            loai, can_cu = "CSDL", "nội dung có bảng trường – kiểu dữ liệu"
        elif not loai and ("ke hoach kiem thu" in k or "pham vi kiem thu" in k):
            loai, can_cu = "Yêu cầu kiểm thử", "nội dung nói về kế hoạch / phạm vi kiểm thử"
    if not loai and duoi in BANG:
        loai, can_cu = "Dữ liệu mẫu", "file bảng tính"
    if duoi in BANG and _du_lieu_that(p):
        canh_bao.append("Có vẻ chứa CCCD/SĐT thật – chỉ dùng cấu trúc và danh mục, không chép giá trị vào kết quả/Artifact/Lark")
    return loai or "Khác", can_cu or "không đoán được – hỏi tester", canh_bao


def _chep(nguon, dich_thu_muc):
    dich = dich_thu_muc / nguon.name
    i = 2
    while dich.exists():
        if dich.stat().st_size == nguon.stat().st_size:
            return dich  # coi như đã chép
        dich = dich_thu_muc / f"{nguon.stem} ({i}){nguon.suffix}"
        i += 1
    shutil.copy2(nguon, dich)
    return dich


def tiep_nhan(da, duong_dan):
    da = Path(da)
    dv = da / "dau-vao"
    dv.mkdir(parents=True, exist_ok=True)
    loi = []
    for s in duong_dan:
        p = Path(s.strip().strip('"'))
        if not p.exists():
            loi.append(f"Không thấy: {s}")
        elif p.is_dir():
            for f in sorted(p.rglob("*")):
                if f.is_file() and f.name.lower() not in BO_QUA and not f.name.startswith("~$"):
                    _chep(f, dv)
        elif p.resolve().parent != dv.resolve():
            _chep(p, dv)
    ds = []
    for f in sorted(dv.iterdir()):
        if f.is_file() and f.name.lower() not in BO_QUA and not f.name.startswith("~$"):
            loai, can_cu, canh_bao = phan_loai(f)
            ds.append({"file": f"dau-vao/{f.name}", "loai": loai, "can_cu": can_cu, "canh_bao": canh_bao})
    theo_loai = {}
    for x in ds:
        theo_loai[x["loai"]] = theo_loai.get(x["loai"], 0) + 1
    thieu = []
    if "SRS" not in theo_loai:
        thieu.append("Tài liệu đặc tả (SRS) – không có thì chỉ kiểm thử thăm dò theo giao diện")
    if "CSDL" not in theo_loai:
        thieu.append("Thiết kế CSDL – dữ liệu sinh theo SRS và giao diện, ràng buộc độ dài/kiểu là giả định")
    if "Checklist thiết lập" not in theo_loai:
        thieu.append("Checklist thiết lập của dev – bước Kiểm tra thiết lập chỉ dựa vào kịch bản")
    kq = {"so_file": len(ds), "theo_loai": theo_loai, "file": ds, "thieu": thieu, "loi": loi}
    (dv / "danh-muc.json").write_text(json.dumps(kq, ensure_ascii=False, indent=2), encoding="utf-8")
    return kq


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    print(json.dumps(tiep_nhan(sys.argv[1], sys.argv[2:]), ensure_ascii=False))


if __name__ == "__main__":
    main()
