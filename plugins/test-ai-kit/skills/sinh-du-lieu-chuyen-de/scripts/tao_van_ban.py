"""Tạo văn bản / hồ sơ mẫu giả lập (tờ khai, đơn, giấy xác nhận, quyết định…) dạng .docx để đính kèm khi test.

Dùng: python tao_van_ban.py <van-ban.json> <thư mục ra>
van-ban.json: danh sách văn bản, mỗi văn bản:
{
  "file": "to-khai-tro-cap-01.docx",
  "quoc_hieu": true,                      // in "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM / Độc lập - Tự do - Hạnh phúc"
  "co_quan": "UBND XÃ THỬ NGHIỆM",        // (tuỳ chọn) góc trái
  "so_ky_hieu": "Số: 01/2026/TK-GL",      // (tuỳ chọn)
  "tieu_de": "TỜ KHAI ĐỀ NGHỊ TRỢ GIÚP XÃ HỘI",
  "phan": [                               // các khối nội dung theo thứ tự
    {"loai": "doan", "chu": "Kính gửi: ..."},
    {"loai": "truong", "nhan": "Họ và tên", "gia_tri": "Nguyễn Văn Thử"},
    {"loai": "bang", "cot": ["STT", "Họ tên"], "dong": [["1", "Trần Thị Mẫu"]]}
  ],
  "ky": {"chuc_danh": "NGƯỜI KHAI", "ten": "Nguyễn Văn Thử"}
}
Mọi văn bản được đóng dấu chữ "DỮ LIỆU GIẢ LẬP – CHỈ DÙNG KIỂM THỬ" ở đầu trang.
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def _doan(tl, chu, dam=False, can=None, co=13, nghieng=False, mau=None):
    p = tl.add_paragraph()
    r = p.add_run(chu)
    r.bold, r.italic, r.font.size = dam, nghieng, Pt(co)
    if mau:
        r.font.color.rgb = RGBColor.from_string(mau)
    if can:
        p.alignment = can
    return p


def tao(vb, thu_muc):
    tl = Document()
    st = tl.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(13)
    _doan(tl, "DỮ LIỆU GIẢ LẬP – CHỈ DÙNG KIỂM THỬ", dam=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=10, mau="C00000")
    if vb.get("co_quan") or vb.get("quoc_hieu", True):
        bang = tl.add_table(rows=1, cols=2)
        trai, phai = bang.rows[0].cells
        trai.text = (vb.get("co_quan", "") + ("\n" + vb["so_ky_hieu"] if vb.get("so_ky_hieu") else "")).strip()
        if vb.get("quoc_hieu", True):
            phai.text = "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc"
        for o in (trai, phai):
            for p in o.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.bold, r.font.size = True, Pt(12)
    _doan(tl, "")
    _doan(tl, vb["tieu_de"].upper(), dam=True, can=WD_ALIGN_PARAGRAPH.CENTER, co=14)
    for k in vb.get("phan", []):
        if k["loai"] == "doan":
            _doan(tl, k["chu"], dam=k.get("dam", False), nghieng=k.get("nghieng", False))
        elif k["loai"] == "truong":
            p = tl.add_paragraph()
            p.add_run(f"{k['nhan']}: ").bold = True
            p.add_run(str(k.get("gia_tri", "")))
        elif k["loai"] == "bang":
            b = tl.add_table(rows=1, cols=len(k["cot"]))
            b.style = "Table Grid"
            for o, c in zip(b.rows[0].cells, k["cot"]):
                o.text = str(c)
                for r in o.paragraphs[0].runs:
                    r.bold = True
            for d in k.get("dong", []):
                for o, gt in zip(b.add_row().cells, d):
                    o.text = str(gt)
    if vb.get("ky"):
        _doan(tl, "")
        _doan(tl, vb["ky"].get("chuc_danh", ""), dam=True, can=WD_ALIGN_PARAGRAPH.RIGHT)
        _doan(tl, "(Ký, ghi rõ họ tên)", nghieng=True, can=WD_ALIGN_PARAGRAPH.RIGHT, co=11)
        _doan(tl, "\n" + vb["ky"].get("ten", ""), dam=True, can=WD_ALIGN_PARAGRAPH.RIGHT)
    duong_dan = Path(thu_muc) / vb["file"]
    tl.save(duong_dan)
    return duong_dan.name


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    ds = json.load(open(sys.argv[1], encoding="utf-8"))
    Path(sys.argv[2]).mkdir(parents=True, exist_ok=True)
    da_tao, loi = [], []
    for vb in ds:
        try:
            da_tao.append(tao(vb, sys.argv[2]))
        except Exception as e:  # noqa: BLE001
            loi.append(f"{vb.get('file', '?')}: {type(e).__name__}: {e}")
    print(json.dumps({"da_tao": da_tao, "loi": loi}, ensure_ascii=False))
    sys.exit(1 if loi else 0)


if __name__ == "__main__":
    main()
